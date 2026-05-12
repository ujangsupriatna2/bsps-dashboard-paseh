#!/usr/bin/env python3
"""
Generate professional BSPS Excel file for Kecamatan Paseh, Desa Loa.
"""

import json
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

# --- Configuration ---
DATA_PATH = "/home/z/my-project/src/data/bsps-data.json"
OUTPUT_PATH = "/home/z/my-project/public/outputs/Data_BSPS_Paseh_Loa.xlsx"

HEADERS = [
    "No", "Nama", "NIK", "No KK", "Alamat", "RT", "RW",
    "Kecamatan", "Desa", "Latitude", "Longitude", "Kategori", "Keterangan"
]

# Column widths (approximate char count)
COL_WIDTHS = [5, 22, 20, 20, 35, 6, 6, 14, 8, 14, 14, 22, 35]

# Colors
HEADER_COLORS = {
    "SEMUA DATA": "4472C4",          # Blue
    "DATA DI ACC": "548235",         # Green
    "TIDAK ACC - LAYAK HUNI": "BF8F00",  # Yellow/Gold
    "TIDAK ACC - TIDAK MELANJUTKAN": "C00000",  # Red
    "RINGKASAN": "1F3864",           # Dark blue
}

ALT_ROW_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Calibri", size=10)
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
SUBTITLE_FONT = Font(name="Calibri", bold=True, size=12)
CATEGORY_FONT = Font(name="Calibri", bold=True, size=11)

CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)


def load_data():
    with open(DATA_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def filter_by_kategori(data, kategori):
    return [d for d in data if d["kategori"] == kategori]


def row_from_entry(entry, no):
    """Convert a JSON entry to a row list."""
    kategori_label = {
        "data_acc": "Data ACC",
        "tidak_acc_layak_huni": "Tidak ACC - Layak Huni",
        "tidak_acc_tidak_melanjutkan": "Tidak ACC - Tidak Melanjutkan",
    }.get(entry["kategori"], entry["kategori"])

    keterangan = entry.get("keterangan") or "-"

    return [
        no,
        entry["nama"],
        str(entry["nik"]),
        str(entry["kk"]),
        entry["alamat"],
        entry["rt"],
        entry["rw"],
        entry["kecamatan"],
        entry["desa"],
        entry["lat"],
        entry["lng"],
        kategori_label,
        keterangan,
    ]


def create_data_sheet(wb, sheet_name, data, header_color):
    """Create a data sheet with formatting."""
    ws = wb.create_sheet(title=sheet_name)

    # Header fill
    header_fill = PatternFill(
        start_color=header_color, end_color=header_color, fill_type="solid"
    )

    # Write headers
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = header_fill
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    # Write data rows
    for row_idx, entry in enumerate(data, 2):
        no = row_idx - 1
        row_data = row_from_entry(entry, no)

        # Alternating row fill
        row_fill = ALT_ROW_FILL if no % 2 == 0 else WHITE_FILL

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.fill = row_fill

            # Alignment
            if col_idx in (1, 6, 7):  # No, RT, RW
                cell.alignment = CENTER_ALIGN
            elif col_idx in (3, 4):  # NIK, No KK - text format, left align
                cell.alignment = LEFT_ALIGN
                cell.number_format = "@"  # Text format
            elif col_idx in (10, 11):  # Lat, Lng
                cell.alignment = CENTER_ALIGN
                cell.number_format = "0.00000000"
            else:
                cell.alignment = LEFT_ALIGN

    # Set column widths
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    # Print settings
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.5)

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(data) + 1}"

    return ws


def create_summary_sheet(wb, all_data, acc_data, layak_huni_data, tidak_melanjutkan_data):
    """Create the RINGKASAN summary sheet."""
    ws = wb.create_sheet(title="RINGKASAN")

    header_fill = PatternFill(
        start_color=HEADER_COLORS["RINGKASAN"],
        end_color=HEADER_COLORS["RINGKASAN"],
        fill_type="solid",
    )

    # --- Title ---
    ws.merge_cells("A1:E1")
    title_cell = ws.cell(row=1, column=1, value="RINGKASAN DATA BSPS KECAMATAN PASEH - DESA LOA")
    title_cell.font = TITLE_FONT
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(1, 6):
        ws.cell(row=1, column=col).fill = header_fill
        ws.cell(row=1, column=col).border = THIN_BORDER

    ws.row_dimensions[1].height = 35

    # --- Summary Stats ---
    row = 3
    stats = [
        ("Total Data", len(all_data)),
        ("Data di ACC", len(acc_data)),
        ("Tidak ACC - Layak Huni", len(layak_huni_data)),
        ("Tidak ACC - Tidak Melanjutkan", len(tidak_melanjutkan_data)),
    ]

    for label, count in stats:
        cell_label = ws.cell(row=row, column=1, value=label)
        cell_label.font = Font(name="Calibri", bold=True, size=11)
        cell_label.alignment = LEFT_ALIGN
        cell_label.border = THIN_BORDER

        cell_value = ws.cell(row=row, column=2, value=count)
        cell_value.font = Font(name="Calibri", bold=True, size=11)
        cell_value.alignment = CENTER_ALIGN
        cell_value.border = THIN_BORDER

        row += 1

    # --- Names by Category ---
    row += 1

    categories = [
        ("DATA DI ACC", acc_data, "548235"),
        ("TIDAK ACC - LAYAK HUNI", layak_huni_data, "BF8F00"),
        ("TIDAK ACC - TIDAK MELANJUTKAN", tidak_melanjutkan_data, "C00000"),
    ]

    for cat_name, cat_data, cat_color in categories:
        # Category header
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cat_cell = ws.cell(row=row, column=1, value=cat_name)
        cat_cell.font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
        cat_fill = PatternFill(start_color=cat_color, end_color=cat_color, fill_type="solid")
        cat_cell.fill = cat_fill
        cat_cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(1, 6):
            ws.cell(row=row, column=col).fill = cat_fill
            ws.cell(row=row, column=col).border = THIN_BORDER
        row += 1

        # Column headers for name list
        sub_headers = ["No", "Nama", "NIK", "Alamat", "Keterangan"]
        for col_idx, hdr in enumerate(sub_headers, 1):
            cell = ws.cell(row=row, column=col_idx, value=hdr)
            cell.font = Font(name="Calibri", bold=True, size=10)
            cell.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
            cell.alignment = CENTER_ALIGN
            cell.border = THIN_BORDER
        row += 1

        # Data rows
        for i, entry in enumerate(cat_data, 1):
            row_fill = ALT_ROW_FILL if i % 2 == 0 else WHITE_FILL

            values = [
                i,
                entry["nama"],
                str(entry["nik"]),
                entry["alamat"],
                entry.get("keterangan") or "-",
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                cell.fill = row_fill
                if col_idx == 1:
                    cell.alignment = CENTER_ALIGN
                elif col_idx == 3:
                    cell.alignment = LEFT_ALIGN
                    cell.number_format = "@"
                else:
                    cell.alignment = LEFT_ALIGN
            row += 1

        row += 1  # Blank row between categories

    # Column widths for summary
    summary_widths = [6, 25, 20, 35, 35]
    for col_idx, width in enumerate(summary_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Print settings
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.5)

    return ws


def main():
    print("Loading BSPS data...")
    all_data = load_data()

    # Filter by category
    acc_data = filter_by_kategori(all_data, "data_acc")
    layak_huni_data = filter_by_kategori(all_data, "tidak_acc_layak_huni")
    tidak_melanjutkan_data = filter_by_kategori(all_data, "tidak_acc_tidak_melanjutkan")

    print(f"Total entries: {len(all_data)}")
    print(f"  Data ACC: {len(acc_data)}")
    print(f"  Tidak ACC - Layak Huni: {len(layak_huni_data)}")
    print(f"  Tidak ACC - Tidak Melanjutkan: {len(tidak_melanjutkan_data)}")

    # Create workbook
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Sheet 1: SEMUA DATA
    print("Creating sheet: SEMUA DATA...")
    create_data_sheet(wb, "SEMUA DATA", all_data, HEADER_COLORS["SEMUA DATA"])

    # Sheet 2: DATA DI ACC
    print("Creating sheet: DATA DI ACC...")
    create_data_sheet(wb, "DATA DI ACC", acc_data, HEADER_COLORS["DATA DI ACC"])

    # Sheet 3: TIDAK ACC - LAYAK HUNI
    print("Creating sheet: TIDAK ACC - LAYAK HUNI...")
    create_data_sheet(
        wb, "TIDAK ACC - LAYAK HUNI", layak_huni_data, HEADER_COLORS["TIDAK ACC - LAYAK HUNI"]
    )

    # Sheet 4: TIDAK ACC - TIDAK MELANJUTKAN
    print("Creating sheet: TIDAK ACC - TIDAK MELANJUTKAN...")
    create_data_sheet(
        wb,
        "TIDAK ACC - TIDAK MELANJUTKAN",
        tidak_melanjutkan_data,
        HEADER_COLORS["TIDAK ACC - TIDAK MELANJUTKAN"],
    )

    # Sheet 5: RINGKASAN
    print("Creating sheet: RINGKASAN...")
    create_summary_sheet(wb, all_data, acc_data, layak_huni_data, tidak_melanjutkan_data)

    # Save
    print(f"Saving to: {OUTPUT_PATH}")
    wb.save(OUTPUT_PATH)
    print("Done! Excel file created successfully.")


if __name__ == "__main__":
    main()
